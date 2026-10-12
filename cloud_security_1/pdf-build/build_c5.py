"""Build Criterion 5 (Cloud Compliance & Governance Controls) evidence PDF + HIPAA checklist workbook."""
import collections
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
src = open(os.path.join(HERE, "build_pdf.py"), encoding="utf-8").read()
exec(src.split("\ns = []\n")[0])  # shared style helpers
from c5_data import AWS_AZURE, CHECKLIST, GUARDRAILS, REMEDIATION, TAG_STANDARD  # noqa: E402

GOV = os.path.join(REPO, "can-governance")
EVID = os.path.join(GOV, "evidence")
OUT = os.path.join(SHOTS, "Care_Access_Navigator_CloudSec1_Criterion_5_Compliance_Governance.pdf")
XLSX = os.path.join(SHOTS, "Care_Access_Navigator_CloudSec1_Criterion_5_HIPAA_Compliance_Checklist.xlsx")


def ev(name):
    return open(os.path.join(EVID, name)).read()


STATUS_COLOR = {"Met": "#1E7B45", "Partial": "#B26A00", "Gap": "#A3271F"}


def status(s):
    return f"<font color='{STATUS_COLOR[s]}'><b>{s}</b></font>"


def c5_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(MUTED)
    canvas.drawString(inch, 0.55 * inch,
                      "Care Access Navigator — CloudSec1 Criterion 5: Cloud Compliance & Governance Controls")
    canvas.drawRightString(letter[0] - inch, 0.55 * inch, f"Page {doc.page}")
    canvas.setStrokeColor(GREY_LINE)
    canvas.line(inch, 0.7 * inch, letter[0] - inch, 0.7 * inch)
    canvas.restoreState()


counts = collections.Counter(r[5] for r in CHECKLIST)
done = sum(1 for r in REMEDIATION if r[6] == "Done")

s = []
# ------------------------------------------------------------------ Cover
s += [P("Care Access Navigator", "title"), Spacer(1, 4),
      P("CloudSec1 Rubric — Criterion 5: Cloud Compliance &amp; Governance Controls", "subtitle"),
      P("Evidence Package — HIPAA Compliance Checklist &amp; Gap Analysis, Management-Group Governance, "
        "Deny Guardrails, Mandatory Tagging and Automated Compliance Checking", "subsub"), Spacer(1, 14)]
s += kv([
    ("Prepared by:", "Akhila Chennamaneni"),
    ("Capstone:", "Care Access Navigator (HealthLine Regional Medical Center)"),
    ("Rubric criterion:", "Criterion 5 — Cloud Compliance &amp; Governance Controls (CloudSec1)"),
    ("Environment:", "Microsoft Azure — tenant chakhilaraogmail.onmicrosoft.com, Azure subscription 1; "
                     "Azure Policy, Management Groups, Defender for Cloud (free tier)"),
    ("Framework:", "HIPAA Security Rule (45 CFR §164.308–§164.312), via the built-in "
                   "HITRUST/HIPAA initiative; NIST CSF 2.0 crosswalk"),
    ("Evidence date:", "October 8, 2026"),
])
s += [P("Objective", "h1"),
      P("This criterion requires mapping cloud configurations to a major compliance framework — with a "
        "checklist, gap analysis and remediation plan — and implementing governance: an organisation "
        "hierarchy with preventive policies, mandatory resource tagging, and automated compliance checking with "
        "documented evidence that the policies are enforced. CAN handles electronic health records, so "
        "<b>HIPAA</b> was chosen as the framework.")]
s += [P("Platform equivalents", "h2"),
      P("The rubric names AWS services. CAN is built on Microsoft Azure, so each control was implemented with "
        "its native Azure equivalent:")]
s += table([["AWS (rubric)", "Azure implementation in CAN"]] + [list(r) for r in AWS_AZURE],
           [2.3 * inch, 4.2 * inch])
s += note("Every governance control in this package uses free Azure features (Management Groups, Azure Policy, "
          "Defender for Cloud foundational CSPM). All enforcement tests either failed before creation or used "
          "validate-only deployments, so testing created no billable resources.", "Cost")

# ------------------------------------------------------------------ Summary
s += [PageBreak(), P("Results at a Glance", "h1")]
s += table([
    ["Requirement", "Implementation", "Verified result"],
    ["Compliance framework mapping", f"HIPAA Security Rule checklist — {len(CHECKLIST)} controls mapped "
     "to CAN's real configuration, with NIST CSF 2.0 crosswalk",
     f"{counts['Met']} Met · {counts['Partial']} Partial · {counts['Gap']} Gap"],
    ["Gap analysis &amp; remediation plan", f"{len(REMEDIATION)} prioritised actions with owner, target date "
     "and cost", f"{done} completed (tagging); remainder scheduled"],
    ["Organisation hierarchy (AWS Organizations)", "mg-can → mg-can-workloads (subscription) / "
     "mg-can-sandbox", ok() + " — subscription moved under governed hierarchy"],
    ["Preventive policies (SCPs)", f"{len(GUARDRAILS)} Deny rules in one initiative at mg-can",
     ok() + " — 7 of 7 guardrail violations blocked"],
    ["Mandatory tagging", "Owner, Environment, DataClassification required (Deny) on resources and RGs",
     ok() + " — untagged resource and RG blocked; 24/24 assets remediated"],
    ["Automated compliance checking (Config Rules)", "HITRUST/HIPAA initiative (586 controls, 90 automated) "
     "assigned audit-only; on-demand scans", ok() + " — continuous evaluation; tags 0 → 72 compliant"],
], [1.8 * inch, 2.6 * inch, 2.1 * inch])

# ------------------------------------------------------------------ Part A governance
s += [P("Part A — Governance Hierarchy (AWS Organizations equivalent)", "h1"),
      P("Azure <b>management groups</b> sit above subscriptions the way AWS Organizations OUs sit above "
        "accounts. Policies assigned to a management group are inherited by every subscription, resource group "
        "and resource beneath it and cannot be removed by a subscription owner — the same property that makes "
        "SCPs effective. CAN's production-style subscription was moved under a dedicated HIPAA workloads group, "
        "with a separate sandbox group that inherits only the organisation guardrails.")]
s += code(ev("00-governance-structure.txt").split("\n== Rules in")[0], "Management groups and policy assignments")
s += optional_shot("Governance - Management groups.png", "mgtree",
                   "Azure portal — Management groups: Care Access Navigator hierarchy with the subscription "
                   "under CAN Workloads (HIPAA).")

# ------------------------------------------------------------------ Part B guardrails
s += step("B1", "Service Control Policy equivalents — organisation guardrails")
s += [P("Seven built-in Azure Policy definitions were combined into one custom initiative, "
        "<b>CAN – Organization guardrails (SCP equivalent)</b>, assigned at the top of the hierarchy "
        "(<font name='Courier'>mg-can</font>) with the <b>Deny</b> effect. Like an SCP, a Deny policy stops the "
        "request at the Azure Resource Manager control plane before anything is created, regardless of the "
        "caller's RBAC role — even a subscription Owner is blocked.")]
s += table([["Rule (reference ID)", "Built-in policy", "Parameter", "Purpose"]] + [list(g) for g in GUARDRAILS],
           [1.45 * inch, 1.75 * inch, 1.6 * inch, 1.7 * inch])
s += [P("A custom non-compliance message explains each denial to the user: <i>“Blocked by CAN guardrails: "
        "only eastus/northcentralus, approved small VM sizes, no public/insecure storage, and no high-cost "
        "network services are permitted.”</i>")]

# ------------------------------------------------------------------ Part C tagging
s += step("C1", "Mandatory resource tagging")
s += [P("The <b>CAN – Mandatory resource tags</b> initiative applies the built-in <i>Require a tag on "
        "resources</i> and <i>Require a tag on resource groups</i> policies once per tag (six rules), with the "
        "Deny effect, at <font name='Courier'>mg-can-workloads</font>. A resource that is missing any of the "
        "three tags cannot be created.")]
s += table([["Tag", "Purpose", "Values"]] + [list(t) for t in TAG_STANDARD], [1.4 * inch, 2.9 * inch, 2.2 * inch])
s += [P("<b>Remediation of existing assets.</b> Policies only block new or changed resources, so the first "
        "scan reported all 24 existing resources and resource groups as non-compliant. All 24 were then tagged; "
        "the six resources that run or store the EHR workload were classified <b>Confidential</b>, network and "
        "platform resources <b>Internal</b>.")]
s += code(ev("04-tag-remediation.txt"), "evidence/04-tag-remediation.txt")

# ------------------------------------------------------------------ Part D compliance checking
s += step("D1", "Automated compliance checking (AWS Config Rules equivalent)")
s += [P("Microsoft's built-in <b>HITRUST/HIPAA</b> regulatory initiative (586 policy references across 576 "
        "control groups) was assigned to <font name='Courier'>mg-can-workloads</font>. Of its controls, 496 are "
        "<i>manual</i> attestations (policies, training, procedures) and 90 are automated checks that Azure "
        "re-evaluates continuously and on every resource change — the equivalent of an AWS Config "
        "conformance pack. Defender for Cloud's default <i>Microsoft cloud security benchmark</i> assignment "
        "(<font name='Courier'>SecurityCenterBuiltIn</font>) provides a second, independent assessment.")]
s += note("The initiative contains 6 policies that can deploy or modify resources (Guest Configuration "
          "extensions, IaaS antimalware, Network Watcher). It was therefore assigned with "
          "<b>enforcementMode = DoNotEnforce</b>: it evaluates and reports on all checks but never changes or "
          "deploys anything, so it cannot create unplanned cost. Enforcement is handled by the explicit Deny "
          "initiatives above.", "Design decision")
s += [P("The tagging remediation moved the mandatory-tags assignment from <b>0 compliant / 72 non-compliant</b> to "
        "<b>72 compliant / 0 non-compliant</b>. The HITRUST/HIPAA findings are unchanged because they need the "
        "planned remediations in Part G; they form the automated input to the gap analysis.")]
s += code(ev("05-compliance-before-after.txt"), "evidence/05-compliance-before-after.txt (latest result per rule)")
s += optional_shot("Governance - Policy compliance.png", "polcomp",
                   "Azure portal — Policy → Compliance for the CAN management groups.")

# ------------------------------------------------------------------ Part E enforcement evidence
s += [CondPageBreak(4.5 * inch), P("Part E — Evidence of Policy Enforcement", "h1"),
      P("Each rule was tested in isolation: test requests carried the mandatory tags (so the tag rule could not "
        "mask other rules) and broke exactly one guardrail. A fully compliant request (T10) was included as a "
        "control to prove the policies block only what they should. Oversized-VM and Bastion tests used "
        "<font name='Courier'>az deployment group validate</font>, which runs the same policy check without "
        "creating anything.")]
s += table([
    ["Test", "Request", "Expected", "Result"],
    ["T1", "NSG with no tags", "Deny (tags)", ok("DENIED") + " — 3 tag rules"],
    ["T2", "Resource group with no tags", "Deny (tags)", ok("DENIED") + " — 3 tag rules"],
    ["T3", "Tagged NSG in centralus", "Deny (region)", ok("DENIED") + " — allowed-locations*"],
    ["T4", "Tagged RG in westus", "Deny (region)", ok("DENIED") + " — allowed-rg-locations"],
    ["T5", "Storage with public blob access", "Deny", ok("DENIED") + " — storage-no-public-access"],
    ["T6", "Storage allowing HTTP", "Deny", ok("DENIED") + " — storage-secure-transfer"],
    ["T7", "Storage with TLS 1.0", "Deny", ok("DENIED") + " — storage-min-tls12"],
    ["T8", "VM Standard_D16s_v5", "Deny (size)", ok("DENIED") + " — allowed-vm-skus"],
    ["T9", "Azure Bastion host", "Deny (type)", ok("DENIED") + " — denied-resource-types"],
    ["T10", "Compliant tagged NSG in eastus", "Allow", ok("ALLOWED") + " — control (then deleted)"],
], [0.45 * inch, 2.4 * inch, 1.25 * inch, 2.4 * inch])
s += code(ev("01-enforcement-tests.txt"), "evidence/01-enforcement-tests.txt")
s += note("*T3's first attempt, one minute after assignment, succeeded because Azure Policy assignments take "
          "up to ~30 minutes to propagate to every resource provider. The test NSG was deleted immediately; the "
          "repeat test after propagation was denied by allowed-locations. A West Europe attempt was also "
          "discarded because Azure itself refused the region (capacity), which is not evidence of our policy.",
          "Issue encountered")
s += optional_shot("Governance - Deny message.png", "denymsg",
                   "Portal or CLI showing a request rejected with the CAN non-compliance message.")

# ------------------------------------------------------------------ Part F checklist
s += [CondPageBreak(4.5 * inch), P("Part F — HIPAA Compliance Checklist", "h1"),
      P("Each HIPAA Security Rule safeguard was mapped to how CAN is actually configured. Status was assigned "
        "from verified configuration and Azure Policy results: <b>Met</b> (control in place and evidenced), "
        "<b>Partial</b> (in place with a material weakness) or <b>Gap</b> (missing). The same checklist, with a "
        "NIST CSF 2.0 crosswalk, is provided as an Excel workbook.")]
rows = [["ID", "HIPAA ref.", "Requirement", "CAN configuration", "Status"]]
for cid, sg, cite, req, cfg, st, evd, nist in CHECKLIST:
    rows.append([cid, f"<nobr>{cite}</nobr>" if len(cite) < 24 else cite, req, f"{cfg}<br/><font size=7.5 color='#777777'>Evidence: {evd}</font>", status(st)])
s += table(rows, [0.4 * inch, 1.3 * inch, 1.2 * inch, 3.0 * inch, 0.6 * inch])

# ------------------------------------------------------------------ Part G gap analysis + remediation
s += [CondPageBreak(4.5 * inch), P("Part G — Gap Analysis &amp; Remediation Plan", "h1")]
by_sg = collections.defaultdict(collections.Counter)
for r in CHECKLIST:
    by_sg[r[1]][r[5]] += 1
s += table([["Safeguard", "Met", "Partial", "Gap"]] +
           [[k, str(v["Met"]), str(v["Partial"]), str(v["Gap"])] for k, v in by_sg.items()] +
           [["<b>Total</b>", f"<b>{counts['Met']}</b>", f"<b>{counts['Partial']}</b>", f"<b>{counts['Gap']}</b>"]],
           [2.9 * inch, 1.2 * inch, 1.2 * inch, 1.2 * inch])
s += [P("Key gaps and their risk", "h2")]
s += bullets([
    "<b>Logging and audit (A3, T3)</b> — no central log store or alerts: a misuse of PHI or an "
    "administrative change could go unnoticed. <i>High.</i> Addressed by Criterion 6.",
    "<b>Privileged access (A5)</b> — the only Owner is a guest personal account: no separation of duties and "
    "a single point of failure. <i>High.</i>",
    "<b>Credentials (H3)</b> — a shared registry admin password and storage keys exist and bypass identity "
    "controls. <i>Medium.</i>",
    "<b>Backup and DR (A9, A10)</b> — no recoverable copy of VM state. <i>Medium</i> in this lab (no PHI on "
    "the VM); <i>High</i> for production.",
    "<b>Network exposure (H1)</b> — the evidence storage account and registry accept traffic from any "
    "network (authentication still required). <i>Medium.</i>",
])
rrows = [["Pri.", "Gap", "Action", "Target", "Cost", "Status"]]
for pri, gid, act, owner, tgt, cost, stt in REMEDIATION:
    col = "#1E7B45" if stt == "Done" else ("#555555" if stt.startswith("Deferred") else "#B26A00")
    rrows.append([pri, gid, f"{act} <font size=7.5 color='#777777'>({owner})</font>", tgt, cost,
                  f"<font color='{col}'><b>{stt}</b></font>"])
s += [P("Remediation plan", "h2")]
s += table(rrows, [0.4 * inch, 0.55 * inch, 3.05 * inch, 0.85 * inch, 0.8 * inch, 0.85 * inch])
s += [P("Priorities: <b>P1</b> free and immediate; <b>P2</b> delivered with Criterion 6 monitoring work; "
        "<b>P3</b> production hardening that carries cost and is deferred under the lab budget.", "caption")]

# ------------------------------------------------------------------ issues + status
s += [KeepTogether([P("Implementation Status", "h2"),
      P("<b>Implemented &amp; verified:</b> management-group hierarchy; organisation Deny guardrails (7 rules); "
        "mandatory tagging (6 rules) with all existing assets remediated; continuous HITRUST/HIPAA evaluation; "
        f"HIPAA checklist of {len(CHECKLIST)} controls with gap analysis and a {len(REMEDIATION)}-item "
        "remediation plan. Policy definitions and all test output are stored in "
        "<font name='Courier'>can-governance/</font>.")])]
s += [P("Issues Encountered", "h1")]
s += table([
    ["Issue", "Resolution"],
    ["Management Groups could not be listed (resource provider not registered)",
     "Registered Microsoft.Management; created hierarchy and moved the subscription"],
    ["HITRUST/HIPAA initiative includes deploy/modify policies that could create billable extensions",
     "Assigned audit-only (DoNotEnforce); explicit Deny initiatives provide enforcement"],
    ["Region test passed during the policy propagation window", "Test resource deleted; re-test after "
     "propagation denied"],
    ["West Europe refused by Azure capacity limits, not by policy", "Discarded; re-tested with centralus"],
    ["'Subnets should have an NSG' flagged 5 subnets that do have NSGs", "Verified association by CLI; treated "
     "as stale Defender assessment. Only AzureBastionSubnet (unused) truly lacks one — scheduled for removal"],
    ["Defender for Containers plan found on Standard (free trial, 20 days left)", "No clusters, so ~$0; added "
     "to plan: switch to Free before trial ends"],
], [3.0 * inch, 3.5 * inch])

doc = SimpleDocTemplate(OUT, pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=0.85 * inch,
                        bottomMargin=0.95 * inch,
                        title="CAN CloudSec1 Criterion 5 - Cloud Compliance & Governance Controls",
                        author="Akhila Chennamaneni")
doc.build(s, onFirstPage=c5_footer, onLaterPages=c5_footer)
print("wrote", OUT)

# ------------------------------------------------------------------ workbook
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

wb = Workbook()
hdr_fill, hdr_font = PatternFill("solid", fgColor="1B2A4A"), Font(bold=True, color="FFFFFF")
fills = {"Met": "D9F2E3", "Partial": "FFF0D1", "Gap": "F8D7D3", "Done": "D9F2E3"}


def sheet(ws, header, rows, widths, status_col=None):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = hdr_fill, hdr_font
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if status_col is not None:
            v = row[status_col].value
            if v in fills:
                row[status_col].fill = PatternFill("solid", fgColor=fills[v])
    ws.freeze_panes = "A2"


ws = wb.active
ws.title = "HIPAA Checklist"
sheet(ws, ["ID", "Safeguard", "HIPAA citation", "Requirement", "CAN configuration", "Status", "Evidence",
           "NIST CSF 2.0"], CHECKLIST, [6, 14, 20, 28, 60, 10, 45, 18], status_col=5)
sheet(wb.create_sheet("Remediation Plan"), ["Priority", "Gap IDs", "Action", "Owner", "Target", "Est. cost",
                                            "Status"], REMEDIATION, [9, 10, 70, 14, 13, 16, 16], status_col=6)
sheet(wb.create_sheet("Guardrails"), ["Reference ID", "Built-in policy", "Parameter", "Purpose"], GUARDRAILS,
      [24, 45, 45, 45])
sheet(wb.create_sheet("Tag Standard"), ["Tag", "Purpose", "Values"], TAG_STANDARD, [22, 50, 40])
summ = wb.create_sheet("Summary")
summ.append(["Status", "Count"])
for c in summ[1]:
    c.fill, c.font = hdr_fill, hdr_font
for k in ("Met", "Partial", "Gap"):
    summ.append([k, f'=COUNTIF(\'HIPAA Checklist\'!F:F,"{k}")'])
summ.append(["Total", "=SUM(B2:B4)"])
summ.column_dimensions["A"].width = 12
wb.save(XLSX)
print("wrote", XLSX)
print("optional screenshots not yet supplied:", MISSING or "none")
