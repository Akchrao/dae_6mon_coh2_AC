"""Compliance tracker workbook: remediation tracking (Task 3) + monitoring/records/documents (Task 4)."""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import org

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "Care_Access_Navigator_Compliance1_Tracker.xlsx")
NAVY, WHITE = "1B2A4A", "FFFFFF"
FILL = {"G": PatternFill("solid", fgColor="D9F2E3"), "A": PatternFill("solid", fgColor="FFF0D1"),
        "R": PatternFill("solid", fgColor="F8D7D3"), "C": PatternFill("solid", fgColor="E5E7EB")}

wb = Workbook()


def sheet(name, header, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    ws.append(header)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color=WHITE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------- Dashboard (filled after KPI sheet)
dash = wb.active
dash.title = "Dashboard"

# ---------------------------------------------------------------- KPI monitor
kpi = sheet("KPI Monitor", ["ID", "KPI", "Definition / formula", "Target", "Amber threshold", "Current", "Unit",
                            "Status", "Source", "Frequency", "Owner", "Measurement note"],
            [], [7, 30, 40, 9, 10, 10, 7, 11, 20, 16, 16, 30])
for i, (kid, name, defi, tgt, amber, cur, unit, src, freq, own, note) in enumerate(org.KPIS, start=2):
    kpi.append([kid, name, defi, tgt, amber, cur, unit, None, src, freq, own, note])
    lower_better = unit in ("count", "min")
    if lower_better:
        f = (f'=IF(F{i}="","Pending",IF(F{i}<=D{i},"Green",IF(F{i}<=E{i},"Amber","Red")))')
    else:
        f = (f'=IF(F{i}="","Pending",IF(F{i}>=D{i},"Green",IF(F{i}>=E{i},"Amber","Red")))')
    kpi.cell(i, 8, f)
last_k = len(org.KPIS) + 1
for txt, key in (("Green", "G"), ("Amber", "A"), ("Red", "R"), ("Pending", "C")):
    kpi.conditional_formatting.add(f"H2:H{last_k}", CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                               fill=FILL[key]))

# ---------------------------------------------------------------- Remediation tracker
rem = sheet("Remediation", ["ID", "Finding", "Rating", "Requirement", "Owner", "Target date", "% complete",
                            "RAG status", "Last update", "Evidence", "Notes"],
            [], [7, 44, 9, 26, 20, 12, 11, 12, 12, 11, 30])
for i, (fid, text, rating, req, owner, target, pct, evid) in enumerate(org.FINDINGS, start=2):
    rem.append([fid, text, rating, req, owner, dt.date.fromisoformat(target), pct / 100, None,
                dt.date(2026, 10, 13), evid, ""])
    rem.cell(i, 6).number_format = "yyyy-mm-dd"
    rem.cell(i, 9).number_format = "yyyy-mm-dd"
    rem.cell(i, 7).number_format = "0%"
    rem.cell(i, 8, f'=IF(G{i}>=1,"Closed",IF(TODAY()>F{i},"Red",IF(AND(F{i}-TODAY()<=7,G{i}<0.75),"Amber",'
                   f'"Green")))')
last_r = len(org.FINDINGS) + 1
for txt, key in (("Green", "G"), ("Amber", "A"), ("Red", "R"), ("Closed", "C")):
    rem.conditional_formatting.add(f"H2:H{last_r}", CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                               fill=FILL[key]))
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", showErrorMessage=True,
                    errorTitle="0-100%", error="Enter a value between 0% and 100%")
rem.add_data_validation(dv)
dv.add(f"G2:G{last_r + 20}")
rem.cell(last_r + 2, 1, "Rules:").font = Font(bold=True)
rem.cell(last_r + 3, 1, "Closed = 100% complete and evidence verified by GRC. Red = past target date. Amber = due "
                        "within 7 days and under 75% complete. Owners update weekly; committee reviews monthly.")

# ---------------------------------------------------------------- Evidence register
ev_rows = []
idx = os.path.join(os.path.dirname(OUT), "audit-evidence", "EVIDENCE-INDEX-SHA256.txt")
for line in open(idx).read().splitlines()[1:]:
    p = line.split()
    if len(p) >= 4:
        ev_rows.append([p[-1].split("_")[0].replace(".json", ""), p[-1], p[0], int(p[1]), "2026-10-13",
                        "GRC Analyst", "audit-evidence/"])
sheet("Evidence Register", ["Evidence ID", "File", "SHA-256", "Bytes", "Collected", "Collected by", "Location"],
      ev_rows, [12, 48, 68, 8, 12, 14, 16])

# ---------------------------------------------------------------- Retention schedule
sheet("Retention Schedule", ["Record type", "Regulation / driver", "Retention", "Storage", "Access",
                             "Owner", "Disposal"], org.RETENTION, [40, 30, 30, 28, 24, 18, 22])

# ---------------------------------------------------------------- Document register
docs = sheet("Document Register", ["Doc ID", "Title", "Version", "Owner", "Approver", "Effective", "Next review",
                                   "Distribution", "Classification", "Location"], org.DOCS,
             [18, 44, 8, 16, 18, 12, 16, 26, 14, 22])

# ---------------------------------------------------------------- Dashboard content
dash["A1"] = "Care Access Navigator — Compliance Monitoring Dashboard"
dash["A1"].font = Font(bold=True, size=14, color=NAVY)
dash["A2"] = "Live formulas: statuses recalculate from the KPI Monitor and Remediation sheets (update 'Current' "
dash["A2"].font = Font(italic=True, color="555555")
dash["A3"] = "and '% complete' values). Last data refresh: 2026-10-13."
dash["A3"].font = Font(italic=True, color="555555")
cards = [("KPIs Green", f"=COUNTIF('KPI Monitor'!H2:H{last_k},\"Green\")"),
         ("KPIs Amber", f"=COUNTIF('KPI Monitor'!H2:H{last_k},\"Amber\")"),
         ("KPIs Red", f"=COUNTIF('KPI Monitor'!H2:H{last_k},\"Red\")"),
         ("Open findings", f"=COUNTIF(Remediation!H2:H{last_r},\"<>Closed\")"),
         ("Open High findings", f"=COUNTIFS(Remediation!C2:C{last_r},\"High\",Remediation!H2:H{last_r},\"<>Closed\")"),
         ("Findings overdue (Red)", f"=COUNTIF(Remediation!H2:H{last_r},\"Red\")")]
for i, (label, formula) in enumerate(cards):
    r = 5 + i
    dash.cell(r, 1, label).font = Font(bold=True)
    dash.cell(r, 2, formula).font = Font(bold=True, size=12, color=NAVY)
dash.column_dimensions["A"].width = 28
dash.column_dimensions["B"].width = 12
# chart data (percent KPIs with values)
dash["D4"], dash["E4"], dash["F4"] = "KPI (percent)", "Current", "Target"
for c in ("D4", "E4", "F4"):
    dash[c].font = Font(bold=True)
row = 5
for kid, name, _, tgt, _, cur, unit, *_ in org.KPIS:
    if unit == "%" and cur is not None:
        dash.cell(row, 4, f"{kid} {name}")
        dash.cell(row, 5, f"='KPI Monitor'!F{[k[0] for k in org.KPIS].index(kid) + 2}")
        dash.cell(row, 6, tgt)
        row += 1
dash.column_dimensions["D"].width = 38
ch = BarChart()
ch.type = "bar"
ch.title = "KPI current vs target (%)"
ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
ch.add_data(Reference(dash, min_col=5, max_col=6, min_row=4, max_row=row - 1), titles_from_data=True)
ch.set_categories(Reference(dash, min_col=4, min_row=5, max_row=row - 1))
ch.height, ch.width = 9, 18
dash.add_chart(ch, "D14")

wb.save(OUT)
print("wrote", OUT)
